"""Read PDF / DOCX / MD / TXT / JSON inputs into raw text (or structured passthrough)."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Optional


@dataclass
class LoadedDocument:
    raw_text: str
    source_path: Path
    detected_format: str
    structured: Optional[dict] = None  # populated when input was JSON


def load_text(path: str | Path) -> LoadedDocument:
    p = Path(path).expanduser().resolve()
    if not p.exists():
        raise FileNotFoundError(p)
    suffix = p.suffix.lower()

    if suffix == ".pdf":
        return LoadedDocument(raw_text=_read_pdf(p), source_path=p, detected_format="pdf")
    if suffix == ".docx":
        return LoadedDocument(raw_text=_read_docx(p), source_path=p, detected_format="docx")
    if suffix in {".md", ".markdown", ".txt"}:
        return LoadedDocument(
            raw_text=p.read_text(encoding="utf-8", errors="replace"),
            source_path=p,
            detected_format=suffix.lstrip("."),
        )
    if suffix == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        # Best-effort: also dump as plain text in case downstream wants it.
        return LoadedDocument(
            raw_text=json.dumps(data, ensure_ascii=False, indent=2),
            source_path=p,
            detected_format="json",
            structured=data,
        )

    raise ValueError(f"Unsupported file type: {suffix}. Supported: .pdf, .docx, .md, .txt, .json")


def _read_pdf(p: Path) -> str:
    """Try pdfplumber for layout-aware extraction; fall back to pypdf."""
    text_parts: list[str] = []
    try:
        import pdfplumber

        with pdfplumber.open(p) as pdf:
            for page in pdf.pages:
                t = page.extract_text(layout=False) or ""
                text_parts.append(t)
        text = "\n\n".join(text_parts).strip()
        if text:
            return text
    except Exception:
        pass

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(p))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
        if text:
            return text
    except Exception as e:
        raise RuntimeError(f"Failed to extract text from PDF {p}: {e}") from e

    raise RuntimeError(f"Could not extract text from PDF {p} (empty result)")


def _read_docx(p: Path) -> str:
    from docx import Document

    doc = Document(str(p))
    blocks: list[str] = []
    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt:
            blocks.append(txt)
    # tables (skills grids etc)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)
